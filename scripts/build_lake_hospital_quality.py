#!/usr/bin/env python3
"""
build_lake_hospital_quality.py — Ingest hospital quality & EPSDT data into the lake.

Reads from: data/raw/hospital_general_info.json
             data/raw/hospital_vbp_scores.json
             data/raw/hospital_hrrp.json
             data/raw/EPSDT416StateRpt2024.xlsx
Writes to:  data/lake/

Tables built:
  Facts:
    fact_hospital_rating     — Hospital overall ratings + quality measure counts
    fact_hospital_vbp        — Value-Based Purchasing total performance scores
    fact_hospital_hrrp       — Hospital Readmissions Reduction Program measures
    fact_epsdt               — EPSDT (CMS-416) children's preventive care by state

Usage:
  python3 scripts/build_lake_hospital_quality.py
  python3 scripts/build_lake_hospital_quality.py --dry-run
  python3 scripts/build_lake_hospital_quality.py --only fact_hospital_vbp
"""

import argparse
import json
import sys
import uuid
from datetime import date, datetime
from pathlib import Path

import duckdb

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = PROJECT_ROOT / "data" / "raw"
LAKE_DIR = PROJECT_ROOT / "data" / "lake"
FACT_DIR = LAKE_DIR / "fact"
META_DIR = LAKE_DIR / "metadata"

SNAPSHOT_DATE = date.today().isoformat()
RUN_ID = str(uuid.uuid4())

HOSPITAL_INFO_JSON = RAW_DIR / "hospital_general_info.json"
VBP_JSON = RAW_DIR / "hospital_vbp_scores.json"
HRRP_JSON = RAW_DIR / "hospital_hrrp.json"
EPSDT_XLSX = RAW_DIR / "EPSDT416StateRpt2024.xlsx"

STATE_NAME_TO_CODE = {
    'Alabama':'AL','Alaska':'AK','Arizona':'AZ','Arkansas':'AR','California':'CA',
    'Colorado':'CO','Connecticut':'CT','Delaware':'DE','District of Columbia':'DC',
    'Florida':'FL','Georgia':'GA','Hawaii':'HI','Idaho':'ID','Illinois':'IL',
    'Indiana':'IN','Iowa':'IA','Kansas':'KS','Kentucky':'KY','Louisiana':'LA',
    'Maine':'ME','Maryland':'MD','Massachusetts':'MA','Michigan':'MI','Minnesota':'MN',
    'Mississippi':'MS','Missouri':'MO','Montana':'MT','Nebraska':'NE','Nevada':'NV',
    'New Hampshire':'NH','New Jersey':'NJ','New Mexico':'NM','New York':'NY',
    'North Carolina':'NC','North Dakota':'ND','Ohio':'OH','Oklahoma':'OK','Oregon':'OR',
    'Pennsylvania':'PA','Rhode Island':'RI','South Carolina':'SC','South Dakota':'SD',
    'Tennessee':'TN','Texas':'TX','Utah':'UT','Vermont':'VT','Virginia':'VA',
    'Washington':'WA','West Virginia':'WV','Wisconsin':'WI','Wyoming':'WY',
    'Puerto Rico':'PR','Guam':'GU','Virgin Islands':'VI',
}


def write_parquet(con, table_name: str, out_path: Path, dry_run: bool) -> int:
    count = con.execute(f"SELECT COUNT(*) FROM {table_name}").fetchone()[0]
    if not dry_run and count > 0:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        con.execute(f"COPY {table_name} TO '{out_path}' (FORMAT PARQUET, COMPRESSION ZSTD)")
        size_mb = out_path.stat().st_size / (1024 * 1024)
        print(f"  -> {out_path.relative_to(LAKE_DIR)} ({count:,} rows, {size_mb:.1f} MB)")
    elif dry_run:
        print(f"  [dry-run] {out_path.relative_to(LAKE_DIR)} ({count:,} rows)")
    return count


def _snapshot_path(fact_name: str) -> Path:
    return FACT_DIR / fact_name / f"snapshot={SNAPSHOT_DATE}" / "data.parquet"


# ---------------------------------------------------------------------------
# Hospital General Information (Care Compare ratings)
# ---------------------------------------------------------------------------

def build_fact_hospital_rating(con, dry_run: bool) -> int:
    print("Building fact_hospital_rating...")
    if not HOSPITAL_INFO_JSON.exists():
        print(f"  SKIPPED — {HOSPITAL_INFO_JSON.name} not found")
        return 0

    con.execute(f"""
        CREATE OR REPLACE TABLE _fact_hosp_rating AS
        SELECT
            facility_id AS provider_id,
            facility_name,
            citytown AS city,
            state AS state_code,
            countyparish AS county,
            zip_code,
            hospital_type,
            hospital_ownership AS ownership_type,
            emergency_services = 'Yes' AS has_emergency,
            meets_criteria_for_birthing_friendly_designation = 'Y' AS birthing_friendly,
            TRY_CAST(hospital_overall_rating AS INTEGER) AS overall_rating,
            -- Mortality
            TRY_CAST(mort_group_measure_count AS INTEGER) AS mort_measure_count,
            TRY_CAST(count_of_mort_measures_better AS INTEGER) AS mort_better,
            TRY_CAST(count_of_mort_measures_no_different AS INTEGER) AS mort_same,
            TRY_CAST(count_of_mort_measures_worse AS INTEGER) AS mort_worse,
            -- Safety
            TRY_CAST(safety_group_measure_count AS INTEGER) AS safety_measure_count,
            TRY_CAST(count_of_safety_measures_better AS INTEGER) AS safety_better,
            TRY_CAST(count_of_safety_measures_no_different AS INTEGER) AS safety_same,
            TRY_CAST(count_of_safety_measures_worse AS INTEGER) AS safety_worse,
            -- Readmission
            TRY_CAST(readm_group_measure_count AS INTEGER) AS readm_measure_count,
            TRY_CAST(count_of_readm_measures_better AS INTEGER) AS readm_better,
            TRY_CAST(count_of_readm_measures_no_different AS INTEGER) AS readm_same,
            TRY_CAST(count_of_readm_measures_worse AS INTEGER) AS readm_worse,
            -- Patient experience
            TRY_CAST(pt_exp_group_measure_count AS INTEGER) AS pt_exp_measure_count,
            TRY_CAST(count_of_facility_pt_exp_measures AS INTEGER) AS pt_exp_facility_measures,
            -- Timeliness & effectiveness
            TRY_CAST(te_group_measure_count AS INTEGER) AS te_measure_count,
            TRY_CAST(count_of_facility_te_measures AS INTEGER) AS te_facility_measures,
            'data.cms.gov/care-compare' AS source,
            DATE '{SNAPSHOT_DATE}' AS snapshot_date
        FROM read_json_auto('{HOSPITAL_INFO_JSON}')
        WHERE state IS NOT NULL AND LENGTH(state) = 2
    """)

    count = write_parquet(con, "_fact_hosp_rating", _snapshot_path("hospital_rating"), dry_run)
    states = con.execute("SELECT COUNT(DISTINCT state_code) FROM _fact_hosp_rating").fetchone()[0]
    avg_r = con.execute("SELECT ROUND(AVG(overall_rating), 2) FROM _fact_hosp_rating WHERE overall_rating > 0").fetchone()[0]
    print(f"  {count:,} hospitals, {states} states, avg rating: {avg_r}")
    con.execute("DROP TABLE IF EXISTS _fact_hosp_rating")
    return count


# ---------------------------------------------------------------------------
# Hospital Value-Based Purchasing (VBP) Scores
# ---------------------------------------------------------------------------

def build_fact_hospital_vbp(con, dry_run: bool) -> int:
    print("Building fact_hospital_vbp...")
    if not VBP_JSON.exists():
        print(f"  SKIPPED — {VBP_JSON.name} not found")
        return 0

    con.execute(f"""
        CREATE OR REPLACE TABLE _fact_hosp_vbp AS
        SELECT
            facility_id AS provider_id,
            facility_name,
            state AS state_code,
            TRY_CAST(fiscal_year AS INTEGER) AS fiscal_year,
            TRY_CAST(unweighted_normalized_clinical_outcomes_domain_score AS DOUBLE) AS clinical_outcomes_score,
            TRY_CAST(weighted_normalized_clinical_outcomes_domain_score AS DOUBLE) AS clinical_outcomes_weighted,
            TRY_CAST(unweighted_person_and_community_engagement_domain_score AS DOUBLE) AS engagement_score,
            TRY_CAST(weighted_person_and_community_engagement_domain_score AS DOUBLE) AS engagement_weighted,
            TRY_CAST(unweighted_normalized_safety_domain_score AS DOUBLE) AS safety_score,
            TRY_CAST(weighted_safety_domain_score AS DOUBLE) AS safety_weighted,
            TRY_CAST(unweighted_normalized_efficiency_and_cost_reduction_domain_score AS DOUBLE) AS efficiency_score,
            TRY_CAST(weighted_efficiency_and_cost_reduction_domain_score AS DOUBLE) AS efficiency_weighted,
            TRY_CAST(total_performance_score AS DOUBLE) AS total_performance_score,
            'data.cms.gov/care-compare' AS source,
            DATE '{SNAPSHOT_DATE}' AS snapshot_date
        FROM read_json_auto('{VBP_JSON}')
        WHERE state IS NOT NULL AND LENGTH(state) = 2
    """)

    count = write_parquet(con, "_fact_hosp_vbp", _snapshot_path("hospital_vbp"), dry_run)
    states = con.execute("SELECT COUNT(DISTINCT state_code) FROM _fact_hosp_vbp").fetchone()[0]
    avg_tps = con.execute("SELECT ROUND(AVG(total_performance_score), 1) FROM _fact_hosp_vbp WHERE total_performance_score > 0").fetchone()[0]
    print(f"  {count:,} hospitals, {states} states, avg TPS: {avg_tps}")
    con.execute("DROP TABLE IF EXISTS _fact_hosp_vbp")
    return count


# ---------------------------------------------------------------------------
# Hospital Readmissions Reduction Program (HRRP)
# ---------------------------------------------------------------------------

def build_fact_hospital_hrrp(con, dry_run: bool) -> int:
    print("Building fact_hospital_hrrp...")
    if not HRRP_JSON.exists():
        print(f"  SKIPPED — {HRRP_JSON.name} not found")
        return 0

    con.execute(f"""
        CREATE OR REPLACE TABLE _fact_hosp_hrrp AS
        SELECT
            facility_id AS provider_id,
            facility_name,
            state AS state_code,
            measure_name,
            TRY_CAST(number_of_discharges AS INTEGER) AS discharges,
            TRY_CAST(excess_readmission_ratio AS DOUBLE) AS excess_readmission_ratio,
            TRY_CAST(predicted_readmission_rate AS DOUBLE) AS predicted_rate,
            TRY_CAST(expected_readmission_rate AS DOUBLE) AS expected_rate,
            TRY_CAST(number_of_readmissions AS INTEGER) AS readmissions,
            start_date,
            end_date,
            'data.cms.gov/care-compare' AS source,
            DATE '{SNAPSHOT_DATE}' AS snapshot_date
        FROM read_json_auto('{HRRP_JSON}')
        WHERE state IS NOT NULL AND LENGTH(state) = 2
    """)

    count = write_parquet(con, "_fact_hosp_hrrp", _snapshot_path("hospital_hrrp"), dry_run)
    states = con.execute("SELECT COUNT(DISTINCT state_code) FROM _fact_hosp_hrrp").fetchone()[0]
    measures = con.execute("SELECT COUNT(DISTINCT measure_name) FROM _fact_hosp_hrrp").fetchone()[0]
    avg_err = con.execute("SELECT ROUND(AVG(excess_readmission_ratio), 4) FROM _fact_hosp_hrrp WHERE excess_readmission_ratio > 0").fetchone()[0]
    print(f"  {count:,} rows, {states} states, {measures} measures, avg ERR: {avg_err}")
    con.execute("DROP TABLE IF EXISTS _fact_hosp_hrrp")
    return count


# ---------------------------------------------------------------------------
# EPSDT (CMS-416) — Children's preventive care participation
# ---------------------------------------------------------------------------

def _parse_num(val):
    """Parse a number from EPSDT formatting (commas, whitespace, DS=suppressed)."""
    if val is None:
        return None
    s = str(val).strip()
    if s in ('', 'DS', 'N/A', '-'):
        return None
    try:
        return float(s.replace(',', ''))
    except ValueError:
        return None


def build_fact_epsdt(con, dry_run: bool) -> int:
    print("Building fact_epsdt...")
    if not EPSDT_XLSX.exists():
        print(f"  SKIPPED — {EPSDT_XLSX.name} not found")
        return 0

    import openpyxl
    wb = openpyxl.load_workbook(str(EPSDT_XLSX), read_only=True)

    # Key EPSDT measures to extract (line item -> field name)
    MEASURES = {
        '1a': 'eligible_total',
        '1b': 'eligible_90_days',
        '1c': 'eligible_chip_expansion',
        '5': 'expected_screenings',
        '6': 'total_screens_received',
        '7': 'screening_ratio',
        '8': 'eligible_should_receive_screen',
        '9': 'eligible_receiving_screen',
        '10': 'participant_ratio',
        '11': 'referred_corrective_treatment',
        '12a': 'receiving_any_dental',
        '12b': 'receiving_preventive_dental',
        '12c': 'receiving_dental_treatment',
        '12g': 'receiving_any_preventive_dental_oral',
        '13': 'enrolled_managed_care',
        '14a': 'blood_lead_tests',
    }

    records = []
    for sheet_name in wb.sheetnames:
        # Extract state name from sheet name (e.g. "Alabama 2024")
        parts = sheet_name.rsplit(' ', 1)
        state_name = parts[0]
        state_code = STATE_NAME_TO_CODE.get(state_name)
        if not state_code:
            continue

        sheet = wb[sheet_name]
        rows = list(sheet.rows)

        record = {'state_code': state_code, 'fiscal_year': 2024}

        for row in rows:
            desc = str(row[0].value or '').strip()
            cat = str(row[1].value or '').strip() if len(row) > 1 else ''
            total = row[2].value if len(row) > 2 else None

            if cat != 'Total':
                continue

            # Match line item number
            for prefix, field in MEASURES.items():
                # Match patterns like "1a.", "5.", "12a."
                if desc.startswith(f'{prefix}.') or desc.startswith(f'{prefix} '):
                    record[field] = _parse_num(total)
                    break

        records.append(record)

    wb.close()

    if not records:
        print("  SKIPPED — no state data parsed")
        return 0

    # Load into DuckDB
    con.execute("CREATE OR REPLACE TABLE _fact_epsdt (state_code VARCHAR, fiscal_year INTEGER, " +
                ", ".join(f"{f} DOUBLE" for f in MEASURES.values()) + ")")

    for rec in records:
        vals = [rec.get('state_code'), rec.get('fiscal_year')]
        vals.extend(rec.get(f) for f in MEASURES.values())
        placeholders = ", ".join(["?" for _ in vals])
        con.execute(f"INSERT INTO _fact_epsdt VALUES ({placeholders})", vals)

    # Add metadata columns
    con.execute(f"""
        CREATE OR REPLACE TABLE _fact_epsdt2 AS
        SELECT *,
            CASE WHEN eligible_total > 0 AND total_screens_received > 0
                THEN ROUND(total_screens_received / eligible_total, 4)
            END AS computed_screening_pct,
            CASE WHEN eligible_total > 0 AND receiving_any_dental > 0
                THEN ROUND(receiving_any_dental / eligible_total, 4)
            END AS dental_pct,
            'medicaid.gov/epsdt/cms-416' AS source,
            DATE '{SNAPSHOT_DATE}' AS snapshot_date
        FROM _fact_epsdt
    """)

    count = write_parquet(con, "_fact_epsdt2", _snapshot_path("epsdt"), dry_run)
    states = con.execute("SELECT COUNT(DISTINCT state_code) FROM _fact_epsdt2").fetchone()[0]
    avg_screen = con.execute("SELECT ROUND(AVG(screening_ratio), 3) FROM _fact_epsdt2 WHERE screening_ratio > 0").fetchone()[0]
    avg_dental = con.execute("SELECT ROUND(AVG(dental_pct), 3) FROM _fact_epsdt2 WHERE dental_pct > 0").fetchone()[0]
    print(f"  {count} states, avg screening ratio: {avg_screen}, avg dental %: {avg_dental}")
    con.execute("DROP TABLE IF EXISTS _fact_epsdt")
    con.execute("DROP TABLE IF EXISTS _fact_epsdt2")
    return count


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

ALL_TABLES = {
    "fact_hospital_rating": build_fact_hospital_rating,
    "fact_hospital_vbp": build_fact_hospital_vbp,
    "fact_hospital_hrrp": build_fact_hospital_hrrp,
    "fact_epsdt": build_fact_epsdt,
}


def main():
    parser = argparse.ArgumentParser(description="Ingest hospital quality & EPSDT data into Aradune lake")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--only", type=str, default=None)
    args = parser.parse_args()

    tables = list(ALL_TABLES.keys())
    if args.only:
        tables = [t.strip() for t in args.only.split(",")]

    print(f"Snapshot: {SNAPSHOT_DATE}")
    print(f"Run ID:   {RUN_ID}")
    print(f"Building: {', '.join(tables)}")
    print()

    con = duckdb.connect()
    totals = {}
    for name in tables:
        totals[name] = ALL_TABLES[name](con, args.dry_run)
        print()

    con.close()

    print("=" * 60)
    print("HOSPITAL QUALITY & EPSDT DATA LAKE INGESTION COMPLETE")
    print("=" * 60)
    total_rows = sum(totals.values())
    for name, count in totals.items():
        status = "written" if not args.dry_run else "dry-run"
        print(f"  {name:35s} {count:>12,} rows  [{status}]")
    print(f"  {'TOTAL':35s} {total_rows:>12,} rows")

    if not args.dry_run and total_rows > 0:
        META_DIR.mkdir(parents=True, exist_ok=True)
        manifest = {
            "snapshot_date": SNAPSHOT_DATE,
            "pipeline_run_id": RUN_ID,
            "created_at": datetime.now().isoformat() + "Z",
            "source": "data.cms.gov",
            "tables": {name: {"rows": count} for name, count in totals.items()},
            "total_rows": total_rows,
        }
        manifest_file = META_DIR / f"manifest_hospital_quality_{SNAPSHOT_DATE}.json"
        with open(manifest_file, "w") as f:
            json.dump(manifest, f, indent=2)
        print(f"\n  Manifest: {manifest_file}")

    print("\nDone.")


if __name__ == "__main__":
    main()
